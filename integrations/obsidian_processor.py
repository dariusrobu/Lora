"""Local processor for user-authored Obsidian notes and handwritten images.

This module creates proposals only. It never writes Lora entities directly;
the existing confirmation flow remains the gate for tasks, exams, and memory.
"""

from __future__ import annotations

import asyncio
import errno
import hashlib
import json
import logging
import re
import subprocess
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
import os
from datetime import datetime
from pathlib import Path
from integrations.obsidian_sync import VAULT_PATH

logger = logging.getLogger("lora.obsidian.processor")
STATE_PATH = VAULT_PATH / ".lora" / "processed.json"
ANALYSIS_DIR = VAULT_PATH / ".lora" / "analysis"
IGNORED_PARTS = {".obsidian", ".trash", ".lora", "00 Dashboard", "Templates"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
DOCUMENT_SUFFIXES = {".pdf", ".pptx", ".ppt", ".docx"}
PROCESS_MODEL = os.getenv("OBSIDIAN_PROCESS_MODEL", "llama3.2:3b")
STUDY_INDEX_NAME = "Study Index.md"
INBOX_STATUS_PATH = VAULT_PATH / "07 Notes" / "Inbox" / "_Processing Status.md"
MAX_FILES_PER_RUN = max(1, int(os.getenv("OBSIDIAN_PROCESS_MAX_FILES", "5")))


def _load_state() -> dict[str, str]:
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def _save_state(state: dict[str, str]) -> None:
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def _digest(path: Path, attempts: int = 5) -> str | None:
    """Hash an iCloud-backed file, tolerating transient provider locks."""
    for attempt in range(attempts):
        try:
            return hashlib.sha256(path.read_bytes()).hexdigest()
        except OSError as exc:
            if exc.errno not in {errno.EAGAIN, errno.EDEADLK, 35} or attempt == attempts - 1:
                raise
            time.sleep(0.5 * (attempt + 1))
    return None


def _is_private_marked(text: str) -> bool:
    head = text[:2000].lower()
    return "private: true" in head or "lora_process: false" in head


def _analysis_target(path: Path) -> Path:
    """Keep summaries beside the course/seminar material that produced them."""
    relative = path.relative_to(VAULT_PATH)
    parts = relative.parts
    for index, part in enumerate(parts):
        if part in {"Cursuri", "Seminarii"} and index > 0:
            return VAULT_PATH.joinpath(*parts[: index + 1], "Lora", "Summaries", f"{path.stem}.summary.md")
    return ANALYSIS_DIR / f"{path.stem}-{hashlib.sha1(str(relative).encode('utf-8')).hexdigest()[:8]}.analysis.md"


def _study_frontmatter(path: Path, kind: str, tags: list[str] | None = None) -> str:
    """Build consistent metadata for Dataview and Graph View links."""
    relative = path.relative_to(VAULT_PATH)
    parts = relative.parts
    subject = ""
    class_type = ""
    if "Subjects" in parts:
        index = parts.index("Subjects")
        if len(parts) > index + 1:
            subject = parts[index + 1]
        if "Cursuri" in parts:
            class_type = "curs"
        elif "Seminarii" in parts:
            class_type = "seminar"
    subject_link = ""
    if subject:
        subject_link = f"05 University/Subjects/{subject}/{subject}"
    source_link = str(relative.with_suffix("")).replace(os.sep, "/")
    return (
        "---\n"
        f"type: {kind}\n"
        + (f"tags: {json.dumps(tags, ensure_ascii=False)}\n" if tags else "")
        + "source: lora\n"
        f"subject: {json.dumps(subject, ensure_ascii=False)}\n"
        f"class_type: {class_type or 'study'}\n"
        f"source_file: {json.dumps(str(relative), ensure_ascii=False)}\n"
        f"subject_link: {json.dumps(subject_link, ensure_ascii=False)}\n"
        f"source_link: {json.dumps(source_link, ensure_ascii=False)}\n"
        f"processed_at: {datetime.now().isoformat(timespec='seconds')}\n"
        "---\n\n"
    )


def _study_links(path: Path) -> str:
    relative = path.relative_to(VAULT_PATH)
    parts = relative.parts
    if "Subjects" not in parts:
        return ""
    index = parts.index("Subjects")
    if len(parts) <= index + 1:
        return ""
    subject = parts[index + 1]
    subject_note = f"[[05 University/Subjects/{subject}/{subject}|{subject}]]"
    resource = f"[[{str(relative.with_suffix('')).replace(os.sep, '/') }|Resource original]]"
    return f"**Materie:** {subject_note}  \n**Sursă:** {resource}\n\n"


def _wikilink(path: Path) -> str:
    """Return a vault-relative Obsidian link without leaking absolute paths."""
    relative = path.relative_to(VAULT_PATH).with_suffix("")
    return "[[" + str(relative).replace(os.sep, "/") + "]]"


def _build_subject_index(subject_dir: Path) -> None:
    """Write a compact, deterministic study index for one subject.

    The index is Lora-owned (inside ``Lora/``), so manually written notes and
    the original resources remain untouched.  It is rebuilt on every scan to
    make newly added iPad files visible without duplicate rows.
    """
    sections: list[str] = []
    for section_name in ("Cursuri", "Seminarii"):
        section_dir = subject_dir / section_name
        if not section_dir.is_dir():
            continue
        resources = sorted(
            p for p in section_dir.rglob("*")
            if p.is_file() and p.suffix.lower() in DOCUMENT_SUFFIXES
        )
        summaries = sorted(
            p for p in section_dir.rglob("*.summary.md") if p.is_file()
        )
        flashcards = sorted(
            p for p in section_dir.rglob("*.flashcards.md") if p.is_file()
        )
        notes = sorted(
            p for p in section_dir.rglob("*.md")
            if p.is_file()
            and p.name not in {"README.md", STUDY_INDEX_NAME}
            and ".summary.md" not in p.name
            and ".flashcards.md" not in p.name
            and "Lora" not in p.parts
        )
        if not any((resources, summaries, flashcards, notes)):
            continue

        lines = [f"## {section_name}"]
        for label, files in (
            ("Resources", resources),
            ("Summaries", summaries),
            ("Flashcards", flashcards),
            ("Notițe", notes),
        ):
            if files:
                lines.append(f"### {label}")
                lines.extend(f"- {_wikilink(path)}" for path in files)
        sections.append("\n".join(lines))

    if not sections:
        return
    target = subject_dir / "Lora" / STUDY_INDEX_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    title = subject_dir.name
    content = (
        "---\n"
        "type: study_index\n"
        "source: lora\n"
        f"subject: {json.dumps(title, ensure_ascii=False)}\n"
        f"updated_at: {datetime.now().isoformat(timespec='seconds')}\n"
        "---\n\n"
        f"# {title} — index de învățare\n\n"
        "Resource-urile originale rămân nemodificate.\n\n"
        + "\n\n".join(sections)
        + "\n"
    )
    target.write_text(content, encoding="utf-8")


def build_study_indexes() -> int:
    """Build indexes for all subject folders; return the number updated."""
    subjects_root = VAULT_PATH / "05 University" / "Subjects"
    if not subjects_root.is_dir():
        return 0
    count = 0
    for subject_dir in sorted(p for p in subjects_root.iterdir() if p.is_dir()):
        before = subject_dir / "Lora" / STUDY_INDEX_NAME
        _build_subject_index(subject_dir)
        if before.exists():
            count += 1
    return count


def build_inbox_status() -> None:
    """Create a read-only processing queue for files placed in Notes/Inbox."""
    inbox = VAULT_PATH / "07 Notes" / "Inbox"
    if not inbox.is_dir():
        return
    subjects_root = VAULT_PATH / "05 University" / "Subjects"
    subjects = [p.name for p in subjects_root.iterdir() if p.is_dir()] if subjects_root.is_dir() else []
    supported = IMAGE_SUFFIXES | DOCUMENT_SUFFIXES | {".md", ".canvas"}
    rows: list[str] = []
    for path in sorted(p for p in inbox.iterdir() if p.is_file() and p.name != INBOX_STATUS_PATH.name):
        if path.suffix.lower() not in supported:
            continue
        haystack = path.stem.casefold()
        subject = next((name for name in subjects if name.casefold() in haystack), None)
        detected = subject or "Necesită alegerea materiei"
        if "seminar" in haystack or "seminarii" in haystack:
            class_folder = "Seminarii"
        elif "curs" in haystack or "lecture" in haystack:
            class_folder = "Cursuri"
        else:
            class_folder = "Cursuri sau Seminarii"
        state = "procesat" if _load_state().get(str(path.relative_to(VAULT_PATH))) == _digest(path) else "nou"
        destination = (
            f"05 University/Subjects/{detected}/{class_folder}/Resources"
            if subject
            else "alege materia și tipul"
        )
        rows.append(
            f"- [[{str(path.relative_to(VAULT_PATH).with_suffix('')).replace(os.sep, '/')}]] "
            f"— **{state}** — Materie: **{detected}** — Destinație propusă: `{destination}`"
        )
    target = VAULT_PATH / "07 Notes" / "Inbox" / INBOX_STATUS_PATH.name
    target.write_text(
        "---\n"
        "type: inbox_status\n"
        "source: lora\n"
        f"updated_at: {datetime.now().isoformat(timespec='seconds')}\n"
        "---\n\n"
        "# Inbox — documente pentru procesare\n\n"
        "Pune aici PDF-uri, PPT-uri, DOCX-uri sau imagini. Lora le analizează local; documentele originale nu sunt mutate automat.\n\n"
        + ("\n".join(rows) if rows else "Nu există documente noi în Inbox.\n"),
        encoding="utf-8",
    )


def move_inbox_document(relative_path: str, subject: str, class_type: str) -> str:
    """Move one Inbox resource after an explicit user confirmation."""
    if class_type not in {"Cursuri", "Seminarii"}:
        raise ValueError("class_type trebuie să fie Cursuri sau Seminarii")
    inbox = (VAULT_PATH / relative_path).resolve()
    inbox_root = (VAULT_PATH / "07 Notes" / "Inbox").resolve()
    if not inbox.is_file() or inbox.parent != inbox_root:
        raise ValueError("Documentul trebuie să fie direct în 07 Notes/Inbox")
    subjects_root = (VAULT_PATH / "05 University" / "Subjects").resolve()
    subject_dir = (subjects_root / subject).resolve()
    if subject_dir.parent != subjects_root or not subject_dir.is_dir():
        raise ValueError("Materia nu există în vault")
    destination = subject_dir / class_type / "Resources" / inbox.name
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        raise FileExistsError(f"Destinația există deja: {destination.name}")
    inbox.rename(destination)
    return str(destination.relative_to(VAULT_PATH))


def build_review_dashboard() -> None:
    """Create the daily study review page using Obsidian Dataview queries."""
    dashboard = VAULT_PATH / "00 Dashboard"
    dashboard.mkdir(parents=True, exist_ok=True)
    target = dashboard / "Review.md"
    target.write_text(
        """---
type: dashboard
source: lora
updated_at: {updated}
---

# Review zilnic

[[Home|Home]] · [[University|University]] · [[Today|Astăzi]] · [[Projects|Proiecte]]

> [!info] Cum folosești review-ul
> Alege un summary, parcurge flashcards și notează întrebările pe care nu le-ai știut. Query-urile sunt locale și se actualizează automat.

## Prioritatea de azi

```dataview
TABLE WITHOUT ID file.link AS "Material", file.mtime AS "Actualizat"
FROM "05 University/Subjects"
WHERE contains(file.name, ".summary") OR contains(file.name, ".flashcards")
SORT file.mtime DESC
LIMIT 5
```

## Flashcards

```dataview
TABLE WITHOUT ID file.link AS "Set", subject AS "Materie", class_type AS "Tip", file.mtime AS "Actualizat"
FROM "05 University/Subjects"
WHERE type = "flashcards"
SORT file.mtime DESC
LIMIT 20
```

## Summaries recente

```dataview
TABLE WITHOUT ID file.link AS "Summary", subject AS "Materie", class_type AS "Tip", file.mtime AS "Actualizat"
FROM "05 University/Subjects"
WHERE type = "study_summary"
SORT file.mtime DESC
LIMIT 20
```

## Resources adăugate recent

```dataview
TABLE WITHOUT ID file.link AS "Resource", file.ext AS "Tip", file.mtime AS "Adăugat"
FROM "05 University/Subjects"
WHERE contains(file.path, "Resources")
SORT file.mtime DESC
LIMIT 15
```

## Examene apropiate

```dataview
TABLE WITHOUT ID file.link AS "Examen", subject_name AS "Materie", exam_date AS "Data"
FROM "05 University/_Data/Exams"
WHERE exam_date >= date(today)
SORT exam_date ASC
LIMIT 10
```

## Taskuri de învățare

```tasks
not done
path includes 02 Tasks
description includes învăț
sort by due
```

> Review-ul este informativ. Lora nu creează taskuri sau modifică date fără confirmare.
""".format(updated=datetime.now().isoformat(timespec="seconds")),
        encoding="utf-8",
    )


def enrich_existing_study_notes() -> int:
    """Add current metadata/links to previously generated study notes."""
    root = VAULT_PATH / "05 University" / "Subjects"
    if not root.is_dir():
        return 0
    updated = 0
    for path in root.rglob("*"):
        if not path.is_file() or not (path.name.endswith(".summary.md") or path.name.endswith(".flashcards.md")):
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        if "source: lora" in text and "subject_link:" in text and "**Materie:**" in text:
            continue
        source_match = re.search(r"^- Source: `([^`]+)`", text, flags=re.MULTILINE)
        source_path = VAULT_PATH / source_match.group(1) if source_match else path
        kind = "flashcards" if path.name.endswith(".flashcards.md") else "study_summary"
        body = text
        if body.startswith("---\n"):
            end = body.find("\n---\n", 4)
            body = body[end + 6:].lstrip() if end >= 0 else body
        prefix = _study_frontmatter(source_path, kind) + _study_links(source_path)
        path.write_text(prefix + body, encoding="utf-8")
        updated += 1
    return updated


async def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages[:40]]
    text = "\n\n".join(page for page in pages if page)
    if text:
        return text
    # Scanned PDFs: render locally with Poppler, then use the existing Apple OCR.
    from core.vision import extract_text_ocr

    ocr_parts: list[str] = []
    with tempfile.TemporaryDirectory(prefix="lora-pdf-") as temp_dir:
        for index in range(min(len(reader.pages), 20)):
            output = Path(temp_dir) / f"page-{index}"
            rendered = subprocess.run(
                ["pdftoppm", "-f", str(index + 1), "-l", str(index + 1), "-singlefile", "-png", str(path), str(output)],
                capture_output=True,
                check=False,
            )
            image_path = output.with_suffix(".png")
            if rendered.returncode == 0 and image_path.exists():
                page_text = await extract_text_ocr(image_path.read_bytes())
                if page_text:
                    ocr_parts.append(f"[Pagina {index + 1}]\n{page_text}")
    return "\n\n".join(ocr_parts)


def _extract_presentation_text(path: Path) -> str:
    if path.suffix.lower() == ".ppt":
        return "Format PPT vechi; exportă prezentarea ca PPTX pentru extragere locală completă."
    try:
        from pptx import Presentation
    except ImportError:
        # Keep PPTX support dependency-light by reading the slide XML directly.
        slides = []
        with zipfile.ZipFile(path) as archive:
            names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            for number, name in enumerate(names, start=1):
                root = ET.fromstring(archive.read(name))
                chunks = [node.text.strip() for node in root.iter() if node.tag.endswith("}t") and node.text and node.text.strip()]
                if chunks:
                    slides.append(f"[Slide {number}]\n" + "\n".join(chunks))
        return "\n\n".join(slides)
    presentation = Presentation(str(path))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                chunks.append(shape.text.strip())
        if chunks:
            slides.append(f"[Slide {number}]\n" + "\n".join(chunks))
    return "\n\n".join(slides)


def _extract_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    paragraphs: list[str] = []
    for paragraph in root.iter():
        if not paragraph.tag.endswith("}p"):
            continue
        chunks = [node.text.strip() for node in paragraph.iter() if node.tag.endswith("}t") and node.text and node.text.strip()]
        if chunks:
            paragraphs.append("".join(chunks))
    return "\n\n".join(paragraphs)


async def _extract_document_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return await _extract_pdf_text(path)
    if path.suffix.lower() == ".docx":
        return _extract_docx_text(path)
    return _extract_presentation_text(path)


async def process_vault() -> dict[str, int]:
    """Create local OCR/LLM proposal notes for newly changed user files."""
    state = _load_state()
    changed = 0
    skipped = 0
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)

    candidates = [path for path in VAULT_PATH.rglob("*") if path.is_file()]
    # New Inbox files are the most useful and should not wait behind a large
    # historical vault.  The cap keeps one run from monopolising Ollama.
    candidates.sort(key=lambda item: ("07 Notes/Inbox" not in str(item), str(item)))
    processed_candidates = 0
    for path in candidates:
        relative_parts = path.relative_to(VAULT_PATH).parts if path.is_relative_to(VAULT_PATH) else path.parts
        if not path.is_file() or any(part in IGNORED_PARTS or part == "Lora" for part in relative_parts):
            continue
        if path.suffix.lower() not in IMAGE_SUFFIXES and path.suffix.lower() not in {".md", ".canvas"} and path.suffix.lower() not in DOCUMENT_SUFFIXES:
            continue
        relative = str(path.relative_to(VAULT_PATH))
        try:
            digest = _digest(path)
        except OSError as exc:
            # iCloud can expose a file while its provider still holds a lock.
            # Leave it untouched and retry on the next scheduled pass.
            logger.warning("Skipping temporarily locked Obsidian file %s: %s", relative, exc)
            skipped += 1
            continue
        if digest is None:
            skipped += 1
            continue
        if state.get(relative) == digest:
            continue
        if processed_candidates >= MAX_FILES_PER_RUN:
            continue
        if path.suffix.lower() == ".md":
            text = path.read_text(encoding="utf-8", errors="replace")
            if "source: \"lora\"" in text or _is_private_marked(text):
                state[relative] = digest
                skipped += 1
                continue
            if "compressed-json" in text and "excalidraw-plugin" in text:
                target = _analysis_target(path)
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(
                    f"# Excalidraw Analysis\n\n- Source: `{relative}`\n- Processed: {datetime.now().isoformat(timespec='seconds')}\n\n"
                    "Desenul Excalidraw este salvat în format comprimat. Pentru OCR local, exportă desenul ca PNG și păstrează imaginea în folderul `Resources`. "
                    "Lora nu va trimite datele comprimate către model și nu va interpreta base64 ca text.\n",
                    encoding="utf-8",
                )
                state[relative] = digest
                changed += 1
                continue
            if len(text.strip()) < 40:
                state[relative] = digest
                skipped += 1
                continue
            from integrations.summary_provider import generate_summary

            try:
                summary = await asyncio.wait_for(generate_summary(
                    "Analizează această notiță Obsidian. Returnează un rezumat scurt, concepte, întrebări și acțiuni posibile. Nu executa acțiuni și nu inventa date.",
                    text[:12000], relative, PROCESS_MODEL,
                ), timeout=180)
            except asyncio.TimeoutError:
                summary = "Ollama nu a răspuns în 180 de secunde. Textul original rămâne disponibil pentru o analiză ulterioară."
            output = _study_frontmatter(path, "note_analysis") + f"# Lora analysis\n\n{_study_links(path)}- Source: `{relative}`\n- Processed: {datetime.now().isoformat(timespec='seconds')}\n\n{summary or 'Ollama nu a returnat o analiză.'}\n\n> Acesta este un draft. Confirmarea rămâne obligatorie în Lora.\n"
        elif path.suffix.lower() == ".canvas":
            try:
                canvas = json.loads(path.read_text(encoding="utf-8"))
                text = "\n".join(str(node.get("text", "")) for node in canvas.get("nodes", []) if node.get("type") == "text")
            except (OSError, json.JSONDecodeError):
                text = ""
            if len(text.strip()) < 20:
                state[relative] = digest
                skipped += 1
                continue
            from integrations.summary_provider import generate_summary

            try:
                summary = await asyncio.wait_for(generate_summary(
                    "Analizează textul extras dintr-un Canvas Obsidian. Returnează rezumat, concepte, întrebări și acțiuni posibile. Nu executa acțiuni și nu inventa date.",
                    text[:12000], relative, PROCESS_MODEL,
                ), timeout=180)
            except asyncio.TimeoutError:
                summary = "Ollama nu a răspuns în 180 de secunde. Reia analiza când modelul este disponibil."
            output = f"# Canvas Analysis\n\n- Source: `{relative}`\n- Processed: {datetime.now().isoformat(timespec='seconds')}\n\n{summary or 'Ollama nu a returnat o analiză.'}\n\n> Acesta este un draft. Confirmarea rămâne obligatorie în Lora.\n"
        elif path.suffix.lower() in DOCUMENT_SUFFIXES:
            text = await _extract_document_text(path)
            if len(text.strip()) < 20:
                state[relative] = digest
                skipped += 1
                continue
            from integrations.summary_provider import generate_summary

            section = "curs" if "Cursuri" in relative.split("/") else "seminar" if "Seminarii" in relative.split("/") else "material de studiu"
            try:
                summary = await asyncio.wait_for(generate_summary(
                    f"Analizează acest material de {section} extras dintr-un PDF, PowerPoint sau document Word. Returnează în română un rezumat clar pe capitole, secțiuni sau slide-uri, concepte-cheie, întrebări de verificare și acțiuni de învățare. Nu executa acțiuni și nu inventa informații sau rezultate numerice; dacă o valoare nu este calculată explicit în material, spune că nu este disponibilă, nu o aproxima.",
                    text[:30000], relative, PROCESS_MODEL,
                ), timeout=180)
            except asyncio.TimeoutError:
                summary = "Ollama nu a răspuns în 180 de secunde. Textul a fost extras local, dar rezumatul trebuie reluat."
            except Exception as exc:
                logger.warning("Study summary unavailable for %s: %s", relative, type(exc).__name__)
                summary = "Analiza nu a putut fi generată acum. Textul original a fost extras local; reia procesarea când modelul este disponibil."
            output = _study_frontmatter(path, "study_summary") + f"# Study Material Analysis\n\n{_study_links(path)}- Source: `{relative}`\n- Processed: {datetime.now().isoformat(timespec='seconds')}\n\n{summary or 'Ollama nu a returnat o analiză.'}\n\n> Documentul original rămâne nemodificat. Confirmarea rămâne obligatorie înainte de a crea acțiuni în Lora.\n"
            try:
                from integrations.summary_provider import generate_flashcards

                flashcards = await asyncio.wait_for(
                    generate_flashcards(summary[:3000], relative, PROCESS_MODEL), timeout=120
                )
                if flashcards.strip():
                    cards_target = _analysis_target(path)
                    cards_target = cards_target.with_name(cards_target.name.replace(".summary.md", ".flashcards.md"))
                    cards_target.parent.mkdir(parents=True, exist_ok=True)
                    cards_target.write_text(
                        _study_frontmatter(path, "flashcards", tags=["flashcards"]) + f"# Flashcards\n\n{_study_links(path)}- Source: `{relative}`\n- Tag: `#flashcards`\n\n{flashcards.strip()}\n",
                        encoding="utf-8",
                    )
            except Exception:
                # A missing quota/model must not prevent the summary from being saved.
                pass
        else:
            from core.vision import extract_text_ocr

            ocr = await extract_text_ocr(path.read_bytes())
            output = f"# OCR Analysis\n\n- Source: `{relative}`\n- Processed: {datetime.now().isoformat(timespec='seconds')}\n\n{ocr or 'Nu am putut extrage text.'}\n\n> OCR local; verifică textul înainte de a crea acțiuni.\n"
        target = _analysis_target(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(output, encoding="utf-8")
        state[relative] = digest
        changed += 1
        processed_candidates += 1

    _save_state(state)
    # Keep one discoverable index per subject even when the latest scan only
    # saw unchanged files.
    build_study_indexes()
    build_inbox_status()
    build_review_dashboard()
    return {"processed": changed, "skipped": skipped}


async def process_vault_safe() -> dict[str, int]:
    try:
        return await process_vault()
    except Exception:
        logger.exception("Obsidian note processing failed")
        return {"processed": 0, "skipped": 0}


if __name__ == "__main__":
    import asyncio

    print(asyncio.run(process_vault_safe()))
